# -*- coding: utf-8 -*-
import base64
import io
import logging
import re

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)

# --- Dependencia: openpyxl ---
try:
    from openpyxl import load_workbook, Workbook
    OPENPYXL_OK = True
except Exception:
    OPENPYXL_OK = False


# ----------------------------- helpers generales -----------------------------

def _to_float(v):
    if v is None:
        return 0.0
    try:
        s = str(v).strip().replace(' ', '').replace(',', '.')
        return float(s) if s else 0.0
    except Exception:
        return 0.0


def _to_int(v):
    try:
        return int(float(v))
    except Exception:
        return 0


# --------------------------------- wizard ------------------------------------

class SnvaImportPackagesWizard(models.TransientModel):
    _name = 'snva.import.packages.wizard'
    _description = 'Import Packages for Simulation (XLSX)'

    simulation_id = fields.Many2one(
        'snva.pallet.simulation', string="Simulation",
        required=True, readonly=True
    )
    file_xlsx = fields.Binary(string="Excel file (.xlsx)")  # NO required (para poder descargar plantilla)
    filename = fields.Char(string="Filename")
    sheet_name = fields.Char(string="Sheet name (optional)")
    has_header = fields.Boolean(string="First row has headers", default=True)
    start_row = fields.Integer(string="Start row (1-based)", default=2)

    help_note = fields.Text(
        readonly=True,
        default=lambda self: _(
            "Se validará con el NOMBRE del producto (product.template.name).\n\n"
            "Encabezados soportados (ES/EN y sinónimos):\n"
            "- Producto: 'Producto', 'Producto en Unidad', 'Descripción', 'Item', 'Artículo'\n"
            "- Marca (opcional): 'Marca'\n"
            "- Cantidad: 'Cantidad', 'Unidades', 'Pieces'\n"
            "- Peso (kg): 'Peso Kg.', 'Peso'\n"
            "- Largo (cm): 'Largo (cm)'\n"
            "- Ancho (cm): 'Ancho (cm)'\n"
            "- Altura (cm): 'Altura (cm)'\n"
        )
    )

    # --------------------------- utilidades internas ---------------------------

    def _ensure_openpyxl(self):
        if not OPENPYXL_OK:
            raise UserError(_("Falta instalar el paquete Python 'openpyxl'."))

    @staticmethod
    def _norm_header(s):
        """minúsculas, sin tildes, sin (..), sin puntos/comas, espacios colapsados"""
        import unicodedata, re as _re
        s = str(s or '').strip().lower()
        s = ''.join(ch for ch in unicodedata.normalize('NFD', s) if unicodedata.category(ch) != 'Mn')
        s = _re.sub(r'\(.*?\)', '', s)
        s = s.replace('.', '').replace(',', '')
        s = ' '.join(s.split())
        return s

    # ------------------------------ importación -------------------------------

    def action_import(self):
        self.ensure_one()
        self._ensure_openpyxl()

        if not self.simulation_id:
            raise UserError(_("Primero guarda la simulación y vuelve a intentar."))

        if not self.file_xlsx:
            raise ValidationError(_("Debes adjuntar un archivo .xlsx"))

        # abrir workbook/hoja
        try:
            wb = load_workbook(io.BytesIO(base64.b64decode(self.file_xlsx)), data_only=True)
        except Exception as e:
            raise UserError(_("No se pudo leer el Excel: %s") % e)
        ws = wb[self.sheet_name] if (self.sheet_name and self.sheet_name in wb.sheetnames) else wb.active

        # encabezados
        header_row_idx = 1 if self.has_header else self.start_row
        try:
            headers = [str(c or '') for c in next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))]
        except StopIteration:
            raise UserError(_("El archivo no contiene encabezados."))

        norm_headers = [self._norm_header(h) for h in headers]

        synonyms = {
            "name":      ["producto en unidad", "producto", "descripcion", "descripción", "item", "articulo", "artículo", "name"],
            "qty":       ["cantidad", "unidades", "pieces", "nro bultos", "unidades por caj", "unidades por caja"],
            "weight_kg": ["peso kg", "peso", "weight"],
            "length_cm": ["largo", "longitud", "length"],
            "width_cm":  ["ancho", "width"],
            "height_cm": ["altura", "alto", "height"],
        }

        def _find_idx(keys):
            for k in keys:
                if k in norm_headers:
                    return norm_headers.index(k)
            return None

        idx = {
            "name": _find_idx(synonyms["name"]),
            "qty": _find_idx(synonyms["qty"]),
            "weight_kg": _find_idx(synonyms["weight_kg"]),
            "length_cm": _find_idx(synonyms["length_cm"]),
            "width_cm": _find_idx(synonyms["width_cm"]),
            "height_cm": _find_idx(synonyms["height_cm"]),
        }

        missing = [k for k in ["name", "qty", "weight_kg", "length_cm", "width_cm", "height_cm"] if idx.get(k) is None]
        if missing:
            raise UserError(_("Faltan columnas requeridas: %s\nEncabezados leídos: %s") % (", ".join(missing), headers))

        Product = self.env['product.template'].sudo()

        start = self.start_row if not self.has_header else max(self.start_row, header_row_idx + 1)
        max_r = ws.max_row

        commands = []       # comandos O2M [(0,0, vals)]
        skipped = []        # mensajes por fila
        missing_products = []  # [(row_idx, name)]
        processed = 0
        will_create = 0

        for row_idx in range(start, max_r + 1):
            row = [c for c in next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))]
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            processed += 1

            name = str(row[idx["name"]] or "").strip()
            if not name:
                skipped.append(_("Fila %s: nombre de producto vacío.") % row_idx)
                continue

            # números
            qty       = max(_to_int(row[idx["qty"]]), 1)
            weight_kg = _to_float(row[idx["weight_kg"]])
            length_cm = _to_float(row[idx["length_cm"]])
            width_cm  = _to_float(row[idx["width_cm"]])
            height_cm = _to_float(row[idx["height_cm"]])

            # buscar producto por nombre (exacto y tolerante)
            prod = Product.search([('name', '=', name)], limit=1)
            if not prod:
                prod = Product.search([('name', 'ilike', name)], limit=1)

            if not prod:
                # ⛳ acumular faltantes para informar claramente
                missing_products.append((row_idx, name))
                skipped.append(_("Fila %s: producto '%s' no está registrado.") % (row_idx, name))
                continue

            vals = {
                # IMPORTANT: con comandos O2M NO se pone snva_package_simulation_id aquí
                'snva_package_name': prod.id,
                'snva_package_width': width_cm,
                'snva_package_height': height_cm,
                'snva_package_length': length_cm,
                'snva_package_weight': weight_kg,
                'snva_package_quantity': qty,
            }
            

            commands.append((0, 0, vals))
            will_create += 1

        # Escribir en el O2M (crea ya enlazado a snva_simulation_package_ids)
        if commands:
            self.simulation_id.sudo().write({'snva_simulation_package_ids': commands})

        # Contar total enlazadas (para feedback)
        total_linked = self.env['snva.pallet.simulation.packages'].sudo().search_count([
            ('snva_package_simulation_id', '=', self.simulation_id.id)
        ])

        # Mensaje final detallado
        msg_lines = [
            _("Filas procesadas: %s") % processed,
            _("Paquetes creados: %s") % will_create,
            _("Total líneas en la simulación: %s") % total_linked,
            _("Filas saltadas: %s") % len(skipped),
        ]

        if missing_products:
            msg_lines.append("")
            msg_lines.append(_("Productos NO registrados (%s):") % len(missing_products))
            # lista breve (máx 30 para no saturar)
            for r, n in missing_products[:30]:
                msg_lines.append("- Fila %s → '%s'" % (r, n))
            if len(missing_products) > 30:
                msg_lines.append("...")

        if skipped:
            msg_lines.append("")
            msg_lines.append(_("Detalles:"))
            for s in skipped[:50]:
                msg_lines.append("- " + s)
            if len(skipped) > 50:
                msg_lines.append("...")

        msg = "\n".join(msg_lines)

        # Notificar (warning si hubo faltantes/errores)
        try:
            if missing_products or skipped:
                self.env.user.notify_warning(message=msg, title=_("Resultado de importación"))
            else:
                self.env.user.notify_success(message=msg, title=_("Resultado de importación"))
        except Exception:
            pass

        # Recargar la vista actual para que aparezcan las líneas
        return {'type': 'ir.actions.client', 'tag': 'reload'}


    # --------------------------- plantilla de ejemplo --------------------------

    def action_download_template(self):
        """Plantilla: Producto, Marca(opc), Cantidad, Peso Kg., Largo(cm), Ancho(cm), Altura(cm)"""
        self.ensure_one()
        self._ensure_openpyxl()

        wb = Workbook()
        ws = wb.active
        ws.title = "Packing"
        headers = ["Producto", "Marca", "Cantidad", "Peso Kg.", "Largo (cm)", "Ancho (cm)", "Altura (cm)"]
        ws.append(headers)
        ws.append(["Caja Standard", "ACME", 10, 5.5, 40, 20, 30])
        ws.append(["Producto Demo", "", 5, 2.0, 25, 15, 10])

        out = io.BytesIO()
        wb.save(out); out.seek(0)
        b64 = base64.b64encode(out.read())
        att = self.env['ir.attachment'].create({
            'name': 'plantilla_packing.xlsx',
            'type': 'binary',
            'datas': b64,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s?download=1' % att.id, 'target': 'self'}
